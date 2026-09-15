"""Unit tests for the extractors that don't already have their own test
module (PyMuPDF4LLMExtractor/AutoPdfExtractor -> test_pdf_text.py/
test_pdf_auto.py). Marker (PdfExtractor) and LibreOffice (OpenOfficeExtractor)
are mocked out rather than exercised for real: loading Marker's models needs
asking the user first per CLAUDE.md, and LibreOffice may not even be
installed on a given machine - these tests only need to prove the extractor
wires its dependency's output into a Document correctly, not that the
dependency itself works."""
import subprocess
import sys
import types as py_types
import zipfile
from pathlib import Path
from unittest.mock import MagicMock

import docx as docx_lib
import pandas as pd
import pytest

from src.extract.archive import ArchiveUnpacker
from src.extract.docx import DocxExtractor
from src.extract.epub import EpubExtractor
from src.extract.html import HtmlExtractor
from src.extract.image import ImageExtractor
from src.extract.openoffice import OpenOfficeExtractor
from src.extract.presentation import PresentationExtractor
from src.extract.spreadsheet import SpreadsheetExtractor
from src.extract.universal import MarkItDownExtractor


# ---------------------------------------------------------------------------
# DocxExtractor
# ---------------------------------------------------------------------------

def test_docx_extractor_extracts_text_and_tables(tmp_path):
    docx_path = tmp_path / "notes.docx"
    d = docx_lib.Document()
    d.add_paragraph("Hello from a paragraph.")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "2"
    d.save(str(docx_path))

    extractor = DocxExtractor()
    doc = extractor.extract(str(docx_path), output_dir=str(tmp_path / "output"))

    assert "Hello from a paragraph." in doc.content
    assert "| A | B |" in doc.content
    assert doc.source_file == "notes.docx"
    assert "error" not in doc.metadata


def test_docx_extractor_handles_unreadable_file(tmp_path):
    bad_path = tmp_path / "not_a_docx.docx"
    bad_path.write_text("this is not a real docx file")

    extractor = DocxExtractor()
    doc = extractor.extract(str(bad_path))

    assert doc.content == ""
    assert "error" in doc.metadata


# ---------------------------------------------------------------------------
# HtmlExtractor
# ---------------------------------------------------------------------------

def test_html_extractor_strips_noisy_tags_and_converts_to_markdown(tmp_path):
    html_path = tmp_path / "page.html"
    html_path.write_text(
        "<html><body>"
        "<nav>Site Nav</nav>"
        "<h1>Real Title</h1>"
        "<p>Real content paragraph.</p>"
        "<footer>Site Footer</footer>"
        "</body></html>"
    )

    extractor = HtmlExtractor(strict_content=True)
    doc = extractor.extract(str(html_path), output_dir=str(tmp_path / "output"))

    assert "Real Title" in doc.content
    assert "Real content paragraph." in doc.content
    assert "Site Nav" not in doc.content
    assert "Site Footer" not in doc.content


def test_html_extractor_handles_parse_failure(tmp_path, monkeypatch):
    html_path = tmp_path / "page.html"
    html_path.write_text("<html></html>")

    extractor = HtmlExtractor()
    monkeypatch.setattr(
        "src.extract.html.md",
        MagicMock(side_effect=RuntimeError("boom")),
    )
    doc = extractor.extract(str(html_path), output_dir=str(tmp_path / "output"))

    assert doc.content == ""
    assert "error" in doc.metadata


# ---------------------------------------------------------------------------
# ImageExtractor
# ---------------------------------------------------------------------------

def test_image_extractor_copies_file_and_emits_markdown_link(tmp_path):
    img_path = tmp_path / "figure.png"
    img_path.write_bytes(b"\x89PNG\r\n\x1a\nfakepngdata")
    output_dir = tmp_path / "output"
    output_dir.mkdir()

    extractor = ImageExtractor()
    doc = extractor.extract(str(img_path), output_dir=str(output_dir))

    assert "![figure]" in doc.content
    assert doc.metadata["extractor"]["images_extracted"] == 1
    copied = list((output_dir.parent / "assets" / "figure").glob("*"))
    assert any(p.name == "figure.png" for p in copied)


# ---------------------------------------------------------------------------
# SpreadsheetExtractor
# ---------------------------------------------------------------------------

def test_spreadsheet_extractor_converts_sheets_to_markdown(tmp_path):
    xlsx_path = tmp_path / "data.xlsx"
    with pd.ExcelWriter(xlsx_path) as writer:
        pd.DataFrame({"col_a": [1, 2], "col_b": [3, 4]}).to_excel(writer, sheet_name="Sheet1", index=False)
        pd.DataFrame().to_excel(writer, sheet_name="EmptySheet", index=False)

    extractor = SpreadsheetExtractor(skip_empty_sheets=True)
    doc = extractor.extract(str(xlsx_path))

    assert "Sheet: Sheet1" in doc.content
    assert "col_a" in doc.content
    assert "EmptySheet" not in doc.content


def test_spreadsheet_extractor_handles_unreadable_file(tmp_path):
    bad_path = tmp_path / "not_a_spreadsheet.xlsx"
    bad_path.write_text("nope")

    extractor = SpreadsheetExtractor()
    doc = extractor.extract(str(bad_path))

    assert doc.content == ""
    assert "error" in doc.metadata


# ---------------------------------------------------------------------------
# PresentationExtractor / MarkItDownExtractor (MarkItDown is mocked - these
# extractors are thin adapters around it, not places to re-test MarkItDown
# itself)
# ---------------------------------------------------------------------------

def test_presentation_extractor_uses_markitdown_text_content(tmp_path):
    pptx_path = tmp_path / "slides.pptx"
    pptx_path.write_bytes(b"fake pptx bytes")

    extractor = PresentationExtractor()
    extractor.md = MagicMock()
    extractor.md.convert.return_value = MagicMock(text_content="Slide 1\n\nSpeaker notes here.")

    doc = extractor.extract(str(pptx_path))

    assert doc.content == "Slide 1\n\nSpeaker notes here."
    assert doc.metadata["extractor"]["mode"] == "presentation"


def test_presentation_extractor_handles_conversion_failure(tmp_path):
    pptx_path = tmp_path / "slides.pptx"
    pptx_path.write_bytes(b"fake pptx bytes")

    extractor = PresentationExtractor()
    extractor.md = MagicMock()
    extractor.md.convert.side_effect = RuntimeError("bad file")

    doc = extractor.extract(str(pptx_path))

    assert doc.content == ""
    assert "error" in doc.metadata


def test_markitdown_extractor_uses_markitdown_text_content(tmp_path):
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("plain text content")

    extractor = MarkItDownExtractor()
    extractor.md = MagicMock()
    extractor.md.convert.return_value = MagicMock(text_content="plain text content")

    doc = extractor.extract(str(txt_path))

    assert doc.content == "plain text content"
    assert doc.metadata["extractor"]["mode"] == "universal_text"


# ---------------------------------------------------------------------------
# ArchiveUnpacker
# ---------------------------------------------------------------------------

def test_archive_unpacker_extracts_zip_contents(tmp_path):
    zip_path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.writestr("a.txt", "hello")
        zf.writestr("sub/b.txt", "world")

    output_dir = tmp_path / "output"
    extractor = ArchiveUnpacker()
    doc = extractor.extract(str(zip_path), output_dir=str(output_dir))

    assert doc.metadata["files_unpacked"] == 2
    unpacked_dir = output_dir / "bundle_unpacked"
    assert (unpacked_dir / "a.txt").exists()
    assert (unpacked_dir / "sub" / "b.txt").exists()
    assert "a.txt" in doc.content


# ---------------------------------------------------------------------------
# EpubExtractor (epub2txt itself is mocked - it's a thin adapter)
# ---------------------------------------------------------------------------

def test_epub_extractor_uses_epub2txt_output(tmp_path, monkeypatch):
    epub_path = tmp_path / "book.epub"
    epub_path.write_bytes(b"fake epub bytes")

    monkeypatch.setattr("src.extract.epub.epub2txt", MagicMock(return_value="Chapter One text."))

    extractor = EpubExtractor()
    doc = extractor.extract(str(epub_path))

    assert doc.content == "Chapter One text."
    assert doc.metadata["extractor"]["mode"] == "epub_text"


def test_epub_extractor_raises_on_parse_failure(tmp_path, monkeypatch):
    epub_path = tmp_path / "book.epub"
    epub_path.write_bytes(b"fake epub bytes")

    monkeypatch.setattr(
        "src.extract.epub.epub2txt",
        MagicMock(side_effect=RuntimeError("corrupt epub")),
    )

    extractor = EpubExtractor()
    with pytest.raises(RuntimeError):
        extractor.extract(str(epub_path))


# ---------------------------------------------------------------------------
# OpenOfficeExtractor (LibreOffice subprocess is mocked)
# ---------------------------------------------------------------------------

def test_openoffice_extractor_delegates_to_docx_extractor_after_conversion(tmp_path, monkeypatch):
    odt_path = tmp_path / "doc.odt"
    odt_path.write_bytes(b"fake odt bytes")

    monkeypatch.setattr("shutil.which", lambda name: "/usr/bin/libreoffice")

    def _fake_run(cmd, check, capture_output, text):
        # Simulate LibreOffice dropping a converted .docx into --outdir.
        outdir = cmd[cmd.index("--outdir") + 1]
        d = docx_lib.Document()
        d.add_paragraph("Converted content.")
        d.save(str(Path(outdir) / "doc.docx"))
        return MagicMock(returncode=0)

    monkeypatch.setattr(subprocess, "run", _fake_run)

    extractor = OpenOfficeExtractor()
    doc = extractor.extract(str(odt_path), output_dir=str(tmp_path / "output"))

    assert "Converted content." in doc.content
    assert doc.metadata["extractor"]["mode"] == "openoffice_conversion"


def test_openoffice_extractor_handles_conversion_failure(tmp_path, monkeypatch):
    odt_path = tmp_path / "doc.odt"
    odt_path.write_bytes(b"fake odt bytes")

    monkeypatch.setattr("shutil.which", lambda name: "/usr/bin/libreoffice")
    monkeypatch.setattr(
        subprocess,
        "run",
        MagicMock(side_effect=subprocess.CalledProcessError(1, "libreoffice", stderr="conversion error")),
    )

    extractor = OpenOfficeExtractor()
    doc = extractor.extract(str(odt_path))

    assert doc.content == ""
    assert "error" in doc.metadata


# ---------------------------------------------------------------------------
# PdfExtractor (Marker) - mocked at the module level so instantiating it
# never actually loads Marker's models (loading them needs asking the user
# first, per CLAUDE.md's hardware rules).
# ---------------------------------------------------------------------------

@pytest.fixture
def fake_marker_modules(monkeypatch):
    fake_models_mod = py_types.ModuleType("marker.models")
    fake_models_mod.load_all_models = MagicMock(return_value=["stub-model-list"])

    fake_convert_mod = py_types.ModuleType("marker.convert")
    fake_convert_mod.convert_single_pdf = MagicMock(
        return_value=("# Extracted heading\n\nBody text.", {}, {"page_count": 1})
    )

    monkeypatch.setitem(sys.modules, "marker.models", fake_models_mod)
    monkeypatch.setitem(sys.modules, "marker.convert", fake_convert_mod)
    return fake_models_mod, fake_convert_mod


def test_pdf_extractor_wires_marker_output_into_document(tmp_path, fake_marker_modules):
    from src.extract.pdf import PdfExtractor

    pdf_path = tmp_path / "book.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 fake")

    extractor = PdfExtractor(batch_size=2)
    doc = extractor.extract(str(pdf_path), output_dir=str(tmp_path / "output"))

    assert "Extracted heading" in doc.content
    assert doc.metadata["extractor"]["program"] == "PdfExtractor (Marker 0.3.10)"


def test_pdf_extractor_raises_when_marker_unavailable(monkeypatch, tmp_path):
    real_import = __import__

    def _blocked_import(name, *args, **kwargs):
        if name == "marker.models":
            raise ImportError("marker not installed")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr("builtins.__import__", _blocked_import)

    from src.extract.pdf import PdfExtractor
    extractor = PdfExtractor()

    with pytest.raises(RuntimeError):
        extractor.extract(str(tmp_path / "book.pdf"))
